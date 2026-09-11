import os
import secrets
import hashlib
from datetime import datetime, date
from io import BytesIO

from fastapi import FastAPI, HTTPException, Depends, Header, Query
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware

from database import get_db, init_db, DB_PATH
from models import (
    LoginRequest, ChangePasswordRequest, SettingsUpdateRequest,
    CustomerCreate, CustomerUpdate, MeasurementCreate, OrderCreate
)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")

app = FastAPI(title="The Jojan - Mens Designer Studio")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Password hashing helpers (pbkdf2, no external deps needed)
# ---------------------------------------------------------------------------

def hash_password(password: str, salt: str = None) -> str:
    if salt is None:
        salt = secrets.token_hex(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 100_000)
    return f"{salt}${dk.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, _ = stored.split("$")
    except ValueError:
        return False
    return hash_password(password, salt) == stored


# ---------------------------------------------------------------------------
# Startup: create tables + default admin user
# ---------------------------------------------------------------------------

@app.on_event("startup")
def startup():
    init_db()
    with get_db() as conn:
        row = conn.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
        if not row:
            conn.execute(
                "INSERT INTO users (username, password_hash, display_name, shop_name) VALUES (?, ?, ?, ?)",
                ("admin", hash_password("admin123"), "Admin", "The Jojan"),
            )


# ---------------------------------------------------------------------------
# Auth dependency
# ---------------------------------------------------------------------------

def get_current_user(authorization: str = Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Not authenticated")
    token = authorization.split(" ", 1)[1]
    with get_db() as conn:
        row = conn.execute(
            """SELECT u.id, u.username, u.display_name, u.shop_name
               FROM sessions s JOIN users u ON s.user_id = u.id
               WHERE s.token = ?""",
            (token,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=401, detail="Invalid or expired session")
    return dict(row)


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------

@app.post("/api/auth/login")
def login(payload: LoginRequest):
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM users WHERE username = ?", (payload.username,)
        ).fetchone()
        if not row or not verify_password(payload.password, row["password_hash"]):
            raise HTTPException(status_code=401, detail="Invalid username or password")
        token = secrets.token_hex(32)
        conn.execute(
            "INSERT INTO sessions (token, user_id) VALUES (?, ?)", (token, row["id"])
        )
    return {
        "token": token,
        "user": {
            "username": row["username"],
            "display_name": row["display_name"],
            "shop_name": row["shop_name"],
        },
    }


@app.post("/api/auth/logout")
def logout(authorization: str = Header(default=None)):
    if authorization and authorization.startswith("Bearer "):
        token = authorization.split(" ", 1)[1]
        with get_db() as conn:
            conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
    return {"ok": True}


@app.get("/api/auth/me")
def me(user=Depends(get_current_user)):
    return user


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@app.get("/api/dashboard")
def dashboard(user=Depends(get_current_user)):
    today = date.today().isoformat()
    with get_db() as conn:
        total_customers = conn.execute("SELECT COUNT(*) c FROM customers").fetchone()["c"]
        todays_orders = conn.execute(
            "SELECT COUNT(*) c FROM orders WHERE order_date = ?", (today,)
        ).fetchone()["c"]
        measurements_saved = conn.execute("SELECT COUNT(*) c FROM measurements").fetchone()["c"]
        recent = conn.execute(
            "SELECT id, customer_code, name, mobile, date_added FROM customers ORDER BY id DESC LIMIT 5"
        ).fetchall()
    return {
        "total_customers": total_customers,
        "todays_orders": todays_orders,
        "measurements_saved": measurements_saved,
        "recent_customers": [dict(r) for r in recent],
        "today": today,
    }


# ---------------------------------------------------------------------------
# Customers
# ---------------------------------------------------------------------------

def generate_customer_code(conn) -> str:
    row = conn.execute("SELECT COUNT(*) c FROM customers").fetchone()
    next_num = row["c"] + 1
    while True:
        code = f"TJ{str(next_num).zfill(6)}"
        exists = conn.execute(
            "SELECT id FROM customers WHERE customer_code = ?", (code,)
        ).fetchone()
        if not exists:
            return code
        next_num += 1


@app.get("/api/customers")
def list_customers(search: str = Query(default=""), user=Depends(get_current_user)):
    with get_db() as conn:
        if search:
            like = f"%{search}%"
            rows = conn.execute(
                """SELECT * FROM customers
                   WHERE name LIKE ? OR mobile LIKE ? OR customer_code LIKE ?
                   ORDER BY id DESC""",
                (like, like, like),
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM customers ORDER BY id DESC").fetchall()
    return [dict(r) for r in rows]


@app.get("/api/customers/next-code")
def next_code(user=Depends(get_current_user)):
    with get_db() as conn:
        return {"customer_code": generate_customer_code(conn)}


@app.post("/api/customers")
def create_customer(payload: CustomerCreate, user=Depends(get_current_user)):
    if not payload.name.strip() or not payload.mobile.strip():
        raise HTTPException(status_code=400, detail="Name and mobile are required")
    with get_db() as conn:
        code = generate_customer_code(conn)
        cur = conn.execute(
            """INSERT INTO customers (customer_code, name, mobile, address, date_added)
               VALUES (?, ?, ?, ?, ?)""",
            (code, payload.name.strip(), payload.mobile.strip(), payload.address, payload.date_added),
        )
        new_id = cur.lastrowid
        row = conn.execute("SELECT * FROM customers WHERE id = ?", (new_id,)).fetchone()
    return dict(row)


@app.get("/api/customers/{customer_id}")
def get_customer(customer_id: int, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Customer not found")
        measurements = conn.execute(
            "SELECT * FROM measurements WHERE customer_id = ? ORDER BY id DESC", (customer_id,)
        ).fetchall()
        orders = conn.execute(
            "SELECT * FROM orders WHERE customer_id = ? ORDER BY id DESC", (customer_id,)
        ).fetchall()
    return {
        "customer": dict(row),
        "measurements": [dict(m) for m in measurements],
        "orders": [dict(o) for o in orders],
    }


@app.put("/api/customers/{customer_id}")
def update_customer(customer_id: int, payload: CustomerUpdate, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT id FROM customers WHERE id = ?", (customer_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Customer not found")
        conn.execute(
            "UPDATE customers SET name = ?, mobile = ?, address = ? WHERE id = ?",
            (payload.name.strip(), payload.mobile.strip(), payload.address, customer_id),
        )
        updated = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    return dict(updated)


@app.delete("/api/customers/{customer_id}")
def delete_customer(customer_id: int, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT id FROM customers WHERE id = ?", (customer_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Customer not found")
        conn.execute("DELETE FROM customers WHERE id = ?", (customer_id,))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Measurements
# ---------------------------------------------------------------------------

@app.post("/api/measurements")
def create_measurement(payload: MeasurementCreate, user=Depends(get_current_user)):
    with get_db() as conn:
        cust = conn.execute(
            "SELECT id FROM customers WHERE id = ?", (payload.customer_id,)
        ).fetchone()
        if not cust:
            raise HTTPException(status_code=404, detail="Customer not found")
        cur = conn.execute(
            """INSERT INTO measurements
               (customer_id, shirt_length, shoulder, chest, waist, sleeve_length, bicep, neck,
                pant_length, pant_waist, hip, thigh, crotch, bottom, special_instructions, entry_date)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                payload.customer_id, payload.shirt_length, payload.shoulder, payload.chest,
                payload.waist, payload.sleeve_length, payload.bicep, payload.neck,
                payload.pant_length, payload.pant_waist, payload.hip, payload.thigh,
                payload.crotch, payload.bottom, payload.special_instructions, payload.entry_date,
            ),
        )
        new_id = cur.lastrowid
        row = conn.execute("SELECT * FROM measurements WHERE id = ?", (new_id,)).fetchone()
    return dict(row)


@app.put("/api/measurements/{measurement_id}")
def update_measurement(measurement_id: int, payload: MeasurementCreate, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT id FROM measurements WHERE id = ?", (measurement_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Measurement not found")
        conn.execute(
            """UPDATE measurements SET
               shirt_length=?, shoulder=?, chest=?, waist=?, sleeve_length=?, bicep=?, neck=?,
               pant_length=?, pant_waist=?, hip=?, thigh=?, crotch=?, bottom=?,
               special_instructions=?, entry_date=?
               WHERE id = ?""",
            (
                payload.shirt_length, payload.shoulder, payload.chest, payload.waist,
                payload.sleeve_length, payload.bicep, payload.neck, payload.pant_length,
                payload.pant_waist, payload.hip, payload.thigh, payload.crotch, payload.bottom,
                payload.special_instructions, payload.entry_date, measurement_id,
            ),
        )
        updated = conn.execute("SELECT * FROM measurements WHERE id = ?", (measurement_id,)).fetchone()
    return dict(updated)


@app.delete("/api/measurements/{measurement_id}")
def delete_measurement(measurement_id: int, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT id FROM measurements WHERE id = ?", (measurement_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Measurement not found")
        conn.execute("DELETE FROM measurements WHERE id = ?", (measurement_id,))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Orders
# ---------------------------------------------------------------------------

@app.post("/api/orders")
def create_order(payload: OrderCreate, user=Depends(get_current_user)):
    with get_db() as conn:
        cust = conn.execute("SELECT id FROM customers WHERE id = ?", (payload.customer_id,)).fetchone()
        if not cust:
            raise HTTPException(status_code=404, detail="Customer not found")
        cur = conn.execute(
            "INSERT INTO orders (customer_id, order_date, status, notes) VALUES (?, ?, ?, ?)",
            (payload.customer_id, payload.order_date, payload.status, payload.notes),
        )
        new_id = cur.lastrowid
        row = conn.execute("SELECT * FROM orders WHERE id = ?", (new_id,)).fetchone()
    return dict(row)


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

@app.put("/api/settings")
def update_settings(payload: SettingsUpdateRequest, user=Depends(get_current_user)):
    with get_db() as conn:
        if payload.display_name is not None:
            conn.execute(
                "UPDATE users SET display_name = ? WHERE id = ?", (payload.display_name, user["id"])
            )
        if payload.shop_name is not None:
            conn.execute(
                "UPDATE users SET shop_name = ? WHERE id = ?", (payload.shop_name, user["id"])
            )
        row = conn.execute("SELECT id, username, display_name, shop_name FROM users WHERE id = ?", (user["id"],)).fetchone()
    return dict(row)


@app.post("/api/settings/change-password")
def change_password(payload: ChangePasswordRequest, user=Depends(get_current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user["id"],)).fetchone()
        if not verify_password(payload.current_password, row["password_hash"]):
            raise HTTPException(status_code=400, detail="Current password is incorrect")
        conn.execute(
            "UPDATE users SET password_hash = ? WHERE id = ?",
            (hash_password(payload.new_password), user["id"]),
        )
    return {"ok": True}


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

@app.get("/api/export/excel")
def export_excel(
    mode: str = Query(default="all"),  # all | date_range | custom
    date_from: str = Query(default=""),
    date_to: str = Query(default=""),
    customer_id: int = Query(default=0),
    user=Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    with get_db() as conn:
        if mode == "custom" and customer_id:
            rows = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchall()
        elif mode == "date_range" and date_from and date_to:
            rows = conn.execute(
                "SELECT * FROM customers WHERE date_added BETWEEN ? AND ? ORDER BY id",
                (date_from, date_to),
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM customers ORDER BY id").fetchall()

        data = []
        for c in rows:
            m = conn.execute(
                "SELECT * FROM measurements WHERE customer_id = ? ORDER BY id DESC LIMIT 1",
                (c["id"],),
            ).fetchone()
            data.append({"customer": dict(c), "measurement": dict(m) if m else {}})

    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Measurements"

    gold = "D4A537"
    dark = "1A1A1A"
    header_fill = PatternFill(start_color=gold, end_color=gold, fill_type="solid")
    title_fill = PatternFill(start_color=dark, end_color=dark, fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True, size=14)
    header_font = Font(color="1A1A1A", bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "ID", "Name", "Mobile", "Address", "Date Added",
        "Shirt Length", "Chest", "Waist(Shirt)", "Shoulder", "Sleeve", "Bicep", "Neck",
        "Pant Length", "Waist(Pant)", "Hip", "Thigh", "Knee/Crotch", "Bottom",
        "Special Instructions",
    ]

    ws.merge_cells("A1:S1")
    ws["A1"] = "THE JOJAN - MENS DESIGNER STUDIO — Customer Measurement Report"
    ws["A1"].font = white_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.append([])
    header_row_idx = 3
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    r = header_row_idx + 1
    for item in data:
        c = item["customer"]
        m = item["measurement"]
        values = [
            c["customer_code"], c["name"], c["mobile"], c.get("address") or "", c["date_added"],
            m.get("shirt_length"), m.get("chest"), m.get("waist"), m.get("shoulder"),
            m.get("sleeve_length"), m.get("bicep"), m.get("neck"),
            m.get("pant_length"), m.get("pant_waist"), m.get("hip"), m.get("thigh"),
            m.get("crotch"), m.get("bottom"), m.get("special_instructions") or "",
        ]
        for col_idx, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=v)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        r += 1

    widths = [10, 18, 14, 20, 12, 11, 9, 11, 10, 9, 8, 8, 11, 11, 8, 8, 12, 9, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Jojan_Customer_Measurements_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Static frontend
# ---------------------------------------------------------------------------

app.mount("/assets", StaticFiles(directory=os.path.join(FRONTEND_DIR, "assets")), name="assets")
app.mount("/css", StaticFiles(directory=os.path.join(FRONTEND_DIR, "css")), name="css")
app.mount("/js", StaticFiles(directory=os.path.join(FRONTEND_DIR, "js")), name="js")


@app.get("/{page_name}.html")
def serve_page(page_name: str):
    path = os.path.join(FRONTEND_DIR, f"{page_name}.html")
    if os.path.exists(path):
        return FileResponse(path)
    raise HTTPException(status_code=404, detail="Page not found")


@app.get("/")
def serve_index():
    return FileResponse(os.path.join(FRONTEND_DIR, "index.html"))
